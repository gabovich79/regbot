from docx import Document

from services.document_service import _docx_body_text
from services.document_structure import boundaries
from services.knowledge import SECTION, prepare_document


def test_contents_entries_do_not_become_body_sections():
    text = 'פתיחה\nתוכן:\n1 . זכויות\n2 . מסלולים\n3 . חריגים\n\nהוראות\nגוף המסמך\n1. הוראה אמיתית\nכלל מחייב'
    positions = boundaries(text, SECTION)
    assert not any('3 . חריגים' in label for label in positions.values())
    assert '1. הוראה אמיתית' in positions.values()
    _,card,_,chunks = prepare_document(text,{'id':1,'title':'מסמך'})
    assert any(s['section']=='תוכן עניינים (ניווט בלבד)' for s in card['section_map'])
    assert any('3 . חריגים' in c['content'] for c in chunks)
    assert all('חריגים' not in c['section'] for c in chunks if 'גוף המסמך' in c['content'])


def test_numbered_body_without_contents_marker_stays_numbered():
    text='1. זכויות\n2. מסלולים\n3. חריגים\n'
    assert list(boundaries(text,SECTION).values()) == ['1. זכויות','2. מסלולים','3. חריגים']


def test_each_table_has_own_locator_and_numbered_cells_stay_in_table():
    doc=Document()
    doc.add_paragraph('תוכן:')
    for label in ['1. הקדמה','2. כללים','3. סיום']: doc.add_paragraph(label)
    doc.add_paragraph('הוראות')
    for title in ['מסלול ראשון','מסלול שני']:
        table=doc.add_table(rows=2,cols=1)
        table.cell(0,0).text=title
        table.cell(1,0).text='1. תנאי בתוך תא\nפרטי הכלל'
    doc.add_paragraph('טקסט אחרי הטבלאות')
    text=_docx_body_text(doc)
    _,card,_,chunks=prepare_document(text,{'id':1,'title':'מסמך'})
    tables=[s['section'] for s in card['section_map'] if s['section'].startswith('טבלה')]
    assert len(tables)==2 and 'מסלול ראשון' in tables[0] and 'מסלול שני' in tables[1]
    assert all(c['section'].startswith('טבלה') for c in chunks if 'פרטי הכלל' in c['content'])
    assert any(c['section']=='מבוא / המשך' and 'טקסט אחרי' in c['content'] for c in chunks)


def test_nested_table_does_not_end_outer_table_early():
    doc=Document(); table=doc.add_table(rows=1,cols=1)
    cell=table.cell(0,0); cell.text='outer heading'
    cell.add_table(rows=1,cols=1).cell(0,0).text='nested evidence'
    cell.add_paragraph('outer continuation')
    text=_docx_body_text(doc)
    positions=boundaries(text,SECTION)
    assert sum(v.startswith('טבלה') for v in positions.values())==1
    _,_,_,chunks=prepare_document(text,{'id':1,'title':'מסמך'})
    assert any('nested evidence' in c['section_text'] and 'outer continuation' in c['section_text'] for c in chunks)

