import io
import json
import sqlite3
from types import SimpleNamespace

import pytest
import httpx
from docx import Document

from scripts.backup_restore import backup,restore
from services.document_service import extract_docx_bytes
from services.acceptance import release_gate,gold_ready
from services.providers import Gateway, BudgetExceeded


def test_docx_table_keeps_document_order():
    doc=Document()
    doc.add_paragraph('לפני הטבלה')
    table=doc.add_table(rows=1, cols=2)
    table.cell(0,0).text='תנאי';table.cell(0,1).text='חריג'
    doc.add_paragraph('אחרי הטבלה')
    blob=io.BytesIO();doc.save(blob)
    text=extract_docx_bytes(blob.getvalue())
    assert text.index('לפני') < text.index('תנאי') < text.index('חריג') < text.index('אחרי')


def test_backup_restore_with_checksums_and_relocated_paths(tmp_path):
    data=tmp_path/'live';data.mkdir()
    text=data/'original.txt';text.write_text('ראיה מקורית',encoding='utf-8')
    db=sqlite3.connect(data/'regbot.db')
    db.execute('CREATE TABLE documents(id INTEGER PRIMARY KEY,text_path TEXT,original_path TEXT)')
    db.execute('INSERT INTO documents VALUES(1,?,?)',(str(text),str(text)))
    db.commit();db.close()
    snapshot=tmp_path/'snapshot'
    backup(data,snapshot)
    restored=tmp_path/'restored'
    restore(snapshot,restored)
    db=sqlite3.connect(restored/'regbot.db')
    path=db.execute('SELECT text_path FROM documents').fetchone()[0]
    db.close()
    from pathlib import Path
    assert Path(path).is_relative_to(restored)
    assert Path(path).read_text(encoding='utf-8')=='ראיה מקורית'
    with pytest.raises(ValueError): restore(snapshot,restored)
    (snapshot/'files/1-text_path.txt').write_text('tampered')
    with pytest.raises(ValueError,match='checksum'): restore(snapshot,tmp_path/'bad')


def test_release_gate_blocks_unreviewed_missing_and_critical_errors():
    assert not gold_ready({'review_status':'pending_source_snapshot'})
    assert not release_gate([])['release_ready']
    cases=[{'id':str(i),'gold_ready':True,'passed':True,'critical_error':False,'required_units':1,'retrieved_units':1,
            'checks':{'correct':True,'complete':True,'supported':True,'handles_missing':True,'handles_conflicts':True},'response_time_ms':1000} for i in range(20)]
    runs=[{'run_id':str(i),'runtime_fingerprint':'runtime','split':'acceptance','case_fingerprint':'same','cases':cases} for i in range(3)]
    assert release_gate(runs)['blockers']==['human_approval_pending']
    assert release_gate(runs,True)['release_ready']
    assert 'three_distinct_runs_required' in release_gate([runs[0]]*3,True)['blockers']
    changed=[dict(r) for r in runs]
    changed[1]['runtime_fingerprint']='different index'
    assert 'runs_must_use_identical_runtime' in release_gate(changed,True)['blockers']
    cases[0]['critical_error']=True
    assert not release_gate(runs,True)['release_ready']


def test_provider_budget_reserves_before_network_and_rejects_unknown_prices(monkeypatch):
    monkeypatch.setenv('PROVIDER_PRICES_JSON','{}')
    gateway=Gateway(limit=.01)
    with pytest.raises(BudgetExceeded):gateway.authorize('model','x',100)
    monkeypatch.setenv('PROVIDER_PRICES_JSON',json.dumps({'model':{'input':100,'output':100}}))
    with pytest.raises(BudgetExceeded):gateway.authorize('model','x',10000)
    assert gateway.spent==0
    with pytest.raises(BudgetExceeded):gateway.authorize('model','x',1,search=True)


@pytest.mark.asyncio
async def test_fetch_rejects_private_dns_before_http(monkeypatch):
    from services import web_evidence
    monkeypatch.setattr(web_evidence.socket,'getaddrinfo',lambda *a,**kw:[(2,1,6,'',('127.0.0.1',443))])
    with pytest.raises(ValueError,match='non-global'):
        await web_evidence.fetch_source('https://www.gov.il/a')


@pytest.mark.asyncio
async def test_fetch_pins_vetted_ip_and_rechecks_redirect(monkeypatch):
    from services import web_evidence
    original=httpx.AsyncClient
    calls=[]
    def handler(request):
        calls.append(request)
        assert request.url.host=='8.8.8.8'
        assert request.headers['host']=='www.gov.il'
        assert request.extensions['sni_hostname']=='www.gov.il'
        return httpx.Response(302,headers={'location':'https://127.0.0.1/private'})
    monkeypatch.setattr(web_evidence.socket,'getaddrinfo',lambda *a,**kw:[(2,1,6,'',('8.8.8.8',443))])
    monkeypatch.setattr(web_evidence.httpx,'AsyncClient',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))
    with pytest.raises(ValueError):await web_evidence.fetch_source('https://www.gov.il/a')
    assert len(calls)==1


@pytest.mark.asyncio
async def test_fetch_uses_fetched_text_not_search_prose(monkeypatch):
    from services import web_evidence
    original=httpx.AsyncClient
    monkeypatch.setattr(web_evidence.socket,'getaddrinfo',lambda *a,**kw:[(2,1,6,'',('8.8.8.8',443))])
    def handler(request):
        return httpx.Response(200,headers={'content-type':'text/html; charset=utf-8'},content='<title>מקור</title><p>ראיה מקורית</p>'.encode())
    monkeypatch.setattr(web_evidence.httpx,'AsyncClient',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))
    result=await web_evidence.fetch_source('https://www.gov.il/a')
    assert 'ראיה מקורית' in result['pages'][0]['content']
    assert result['source_hash'] and result['retrieved_at']


def test_genai_configuration_supported_by_pinned_sdk():
    from google.genai import types
    value=types.GenerateContentConfig(thinking_config=types.ThinkingConfig(thinking_budget=0),response_mime_type='application/json')
    assert value.thinking_config.thinking_budget==0
    assert types.HttpRetryOptions(attempts=1).attempts==1
