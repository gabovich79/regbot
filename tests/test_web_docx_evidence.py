import io
import zipfile

import httpx
import pytest
from docx import Document

from services import web_evidence


def word_source():
    document = Document()
    document.add_paragraph('before')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'condition'
    table.cell(0, 1).text = 'exception'
    document.add_paragraph('after')
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_web_word_evidence_preserves_table_order(monkeypatch):
    original = httpx.AsyncClient
    monkeypatch.setattr(web_evidence.socket, 'getaddrinfo', lambda *a, **kw: [(2, 1, 6, '', ('8.8.8.8', 443))])
    payload = word_source()
    def handler(request):
        return httpx.Response(200, headers={'content-type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}, content=payload)
    monkeypatch.setattr(web_evidence.httpx, 'AsyncClient', lambda **kw: original(transport=httpx.MockTransport(handler), **kw))
    source = await web_evidence.fetch_source('https://www.gov.il/source.docx')
    text = source['pages'][0]['content']
    assert text.index('before') < text.index('condition') < text.index('exception') < text.index('after')
    assert source['kind'] == 'official_web' and source['source_hash'] and source['retrieved_at']


def test_compressed_oversize_word_rejected_before_parser(monkeypatch):
    monkeypatch.setattr(web_evidence, 'MAX_BYTES', 64)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', 'x'*1000)
    with pytest.raises(ValueError, match='size limit'):
        web_evidence.read_docx_source(buffer.getvalue())


def test_unrelated_zip_is_not_word_evidence():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        archive.writestr('unrelated.txt', 'text')
    with pytest.raises(ValueError, match='Word document'):
        web_evidence.read_docx_source(buffer.getvalue())
