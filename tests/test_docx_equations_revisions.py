import io

import pytest
from docx import Document
from docx.oxml import parse_xml

from services.document_service import extract_docx_bytes
from services.knowledge import quality_issues

M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def extract(doc):
    stream = io.BytesIO()
    doc.save(stream)
    return extract_docx_bytes(stream.getvalue())


def test_equation_fraction_survives_in_body_and_table_in_order():
    doc = Document()
    doc.add_paragraph('before')
    paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph._p.append(parse_xml(f'<m:oMath xmlns:m="{M}"><m:f><m:num><m:r><m:t>d</m:t></m:r></m:num><m:den><m:r><m:t>20</m:t></m:r></m:den></m:f><m:r><m:t>=FC</m:t></m:r></m:oMath>'))
    doc.add_paragraph('after')
    text = extract(doc)
    assert '(d)/(20)=FC' in text
    assert text.index('before') < text.index('(d)/(20)=FC') < text.index('after')
    assert 'd20=FC' not in text


def test_struck_old_date_is_marked_not_silently_accepted_or_deleted():
    doc = Document()
    p = doc.add_paragraph('תחילה: ')
    p.add_run('1 בפברואר 2020').font.strike = True
    p.add_run(' 25 ביוני 2020').font.underline = True
    text = extract(doc)
    assert '[מחוק במקור: 1 בפברואר 2020]' in text
    assert '25 ביוני 2020' in text
    assert 'marked_revisions_require_source_review' in quality_issues(text)


def test_explicit_strike_off_and_whitespace_are_not_deletions():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('current text').font.strike = False
    p.add_run(' ').font.strike = True
    assert extract(doc) == 'current text'


def test_tracked_replacement_retains_both_roles():
    doc = Document()
    doc.add_paragraph()._p.append(parse_xml(f'<w:del xmlns:w="{W}"><w:r><w:delText>old</w:delText></w:r></w:del>'))
    doc.paragraphs[0]._p.append(parse_xml(f'<w:ins xmlns:w="{W}"><w:r><w:t>new</w:t></w:r></w:ins>'))
    text = extract(doc)
    assert '[מחוק במקור: old]' in text
    assert '[תוספת מסומנת במקור: new]' in text


def test_merged_cells_preserve_span_and_continuation_without_repeating_rules():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = 'shared heading'
    table.cell(1, 0).merge(table.cell(2, 0)).text = 'shared condition'
    table.cell(1, 1).text = 'first exception'
    table.cell(2, 1).text = 'second exception'
    text = extract(doc)
    assert text.count('shared heading') == 1
    assert text.count('shared condition') == 1
    assert '[פריסת תא במקור: 3 עמודות]' in text
    assert '[מיזוג אנכי במקור: המשך התא מעל]' in text
    assert text.index('first exception') < text.index('second exception')


def test_nested_table_and_repeated_distinct_cells_are_not_lost():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = 'before nested'
    cell.add_table(rows=1, cols=1).cell(0, 0).text = 'nested requirement'
    cell.add_paragraph('after nested')
    table.cell(0, 1).text = 'nested requirement'
    text = extract(doc)
    assert text.count('nested requirement') == 2
    assert text.index('before nested') < text.index('nested requirement') < text.index('after nested')


@pytest.mark.parametrize('body', [
    '<m:rad><m:e><m:r><m:t>x</m:t></m:r></m:e></m:rad>',
    '<m:f><m:num><m:r><m:t>x</m:t></m:r></m:num></m:f>',
    '<m:f><m:fPr><m:type m:val="noBar"/></m:fPr><m:num/><m:den/></m:f>',
])
def test_unsupported_equation_fails_instead_of_inventing_a_formula(body):
    doc = Document()
    doc.add_paragraph()._p.append(parse_xml(f'<m:oMath xmlns:m="{M}">{body}</m:oMath>'))
    with pytest.raises(ValueError, match='DOCX equation'):
        extract(doc)
