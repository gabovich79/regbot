from docx import Document

from services.document_service import extract_docx, extract_docx_structure
from services.knowledge import prepare_document


def source(tmp_path):
    doc = Document()
    for number in range(1, 4):
        doc.add_paragraph().add_run(f'{number}. כותרת ראשית').bold = True
        doc.add_paragraph(f'1. תנאי פנימי עבור כלל {number}')
        doc.add_paragraph(f'2. חריג פנימי עבור כלל {number}')
    path = tmp_path/'source.docx'
    doc.save(path)
    return path


def test_bold_source_titles_keep_numbered_conditions_inside_parent(tmp_path):
    path = source(tmp_path)
    text, layout = extract_docx_structure(path)
    assert text == extract_docx(path)
    _, card, issues, chunks = prepare_document(text, {'id':1, 'title':'מקור', 'original_path':str(path)})
    assert len(card['section_map']) == 3
    assert 'docx_hierarchy_inferred_from_typography_requires_review' in issues
    assert len(card['structure_review']['suppressed_numbered_boundaries']) == 6
    parents = [chunks[s['first_chunk']]['section_text'] for s in card['section_map']]
    assert ''.join(parents) == text
    assert 'תנאי פנימי עבור כלל 2' in parents[1]
    assert 'חריג פנימי עבור כלל 2' in parents[1]


def test_typography_from_different_text_is_not_applied(tmp_path):
    path = source(tmp_path)
    text = extract_docx(path)+'\n4. קטע שלא נמצא במקור'
    _, card, issues, _ = prepare_document(text, {'id':1, 'title':'מקור', 'original_path':str(path)})
    assert card['structure_review']['mode'] == 'text_patterns'
    assert 'docx_structure_source_mismatch_requires_reextraction' in issues


def test_repeated_nested_paragraphs_are_aligned_in_source_order(tmp_path):
    doc = Document()
    for number in range(1, 4):
        doc.add_paragraph().add_run(f'{number}. כותרת ראשית').bold = True
        doc.add_paragraph('1. תנאי זהה החוזר במספר מקומות')
    path = tmp_path/'repeated.docx'
    doc.save(path)
    text, layout = extract_docx_structure(path)
    assert layout['alignment_complete']
    assert len({p['start'] for p in layout['paragraphs']}) == 6
    _, card, _, chunks = prepare_document(text, {'id':1, 'title':'מקור', 'original_path':str(path)})
    assert len(card['section_map']) == 3
    assert all('תנאי זהה' in chunks[s['first_chunk']]['section_text'] for s in card['section_map'])
