from docx import Document
from docx.oxml import parse_xml
from docx.enum.style import WD_STYLE_TYPE

from services.docx_numbering import Numbering, W, UNRESOLVED
from services.document_service import _docx_body_text
from services.knowledge import quality_issues

NS = W[1:-1]


def definition(doc, kind='decimal', extra=''):
    root = doc.part.numbering_part.element
    root.append(parse_xml(f'<w:abstractNum xmlns:w="{NS}" w:abstractNumId="99">'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="{kind}"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2)"/></w:lvl>'
        '</w:abstractNum>'))
    root.append(parse_xml(f'<w:num xmlns:w="{NS}" w:numId="99"><w:abstractNumId w:val="99"/>{extra}</w:num>'))


def paragraph(doc, text, level=0, ident=99):
    p = doc.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(f'<w:numPr xmlns:w="{NS}"><w:ilvl w:val="{level}"/><w:numId w:val="{ident}"/></w:numPr>'))
    return p


def test_multilevel_restart_and_table_continuation():
    doc = Document(); definition(doc)
    paragraph(doc, 'first'); paragraph(doc, 'child', 1); paragraph(doc, 'child two', 1)
    paragraph(doc, 'second'); paragraph(doc, 'reset child', 1)
    table = doc.add_table(rows=1, cols=1)
    p = paragraph(doc, 'inside table')
    table.cell(0,0)._tc.append(p._p)
    text = _docx_body_text(doc)
    for value in ['1. first','1.1) child','1.2) child two','2. second','2.1) reset child','3. inside table']:
        assert value in text


def test_style_inheritance_override_start_and_explicit_disable():
    doc = Document(); definition(doc, extra='<w:lvlOverride w:ilvl="0"><w:startOverride w:val="7"/></w:lvlOverride>')
    base = doc.styles.add_style('Numbered base', WD_STYLE_TYPE.PARAGRAPH)
    base.element.get_or_add_pPr().append(parse_xml(f'<w:numPr xmlns:w="{NS}"><w:numId w:val="99"/></w:numPr>'))
    derived = doc.styles.add_style('Numbered child', WD_STYLE_TYPE.PARAGRAPH); derived.base_style = base
    doc.add_paragraph('inherited', derived)
    p = paragraph(doc, 'disabled', ident=0); p.style = derived
    doc.add_paragraph('next', derived)
    text = _docx_body_text(doc)
    assert '7. inherited' in text and '8. next' in text and '\ndisabled\n' in text


def test_unsupported_format_is_visible_and_flagged():
    doc = Document(); definition(doc, kind='unknownFormat'); paragraph(doc,'rule')
    text = _docx_body_text(doc)
    assert UNRESOLVED in text
    assert 'unresolved_source_numbering_requires_review' in quality_issues(text)


def test_hebrew_labels_and_independent_extractions():
    doc = Document(); definition(doc, kind='hebrew1')
    paragraph(doc, 'first'); paragraph(doc, 'second')
    assert _docx_body_text(doc) == 'א. first\nב. second'
    assert _docx_body_text(doc) == 'א. first\nב. second'
